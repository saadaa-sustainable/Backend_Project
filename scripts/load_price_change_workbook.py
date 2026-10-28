"""Turn a price-change workbook into the JSON the impact loader reads.

Kept apart from refresh_sku_price_impact.py so the workbook — which is a
commercial document that arrives by hand — is converted once, explicitly,
rather than being re-read on every refresh.

The output lands in scripts/data/, which is gitignored: this repo is
PUBLIC and a full price list across the catalogue is not something to
commit. The prices themselves live in Supabase.

Expects two sheets, as the 2026-10-28 workbook had them:
    "Price Changes"  Master SKU | Current Price | Updated Price | Change | Status
    "SKUs to Hide"   one SKU per row

Usage:
    ./.venv/bin/python scripts/load_price_change_workbook.py ~/Downloads/SAADAA_Price_Changes.xlsx \
        --effective 2026-10-28
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--effective", required=True,
                    help="IST date the new prices went live (YYYY-MM-DD). "
                         "Confirm it against each product's Shopify updatedAt "
                         "rather than assuming the day the file was sent.")
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    changes = [
        {"master_sku": r[0], "old_price": r[1], "new_price": r[2],
         "change": r[3], "status": r[4]}
        for r in wb["Price Changes"].iter_rows(min_row=2, values_only=True) if r[0]
    ]
    hide = [r[0] for r in wb["SKUs to Hide"].iter_rows(min_row=2, values_only=True) if r[0]]

    out = args.out or (Path(__file__).parent / "data" /
                       f"sku_price_changes_{args.effective.replace('-', '_')}.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(
        {"effective_date": args.effective, "changes": changes, "hide": hide}, indent=2))

    moved = sum(1 for c in changes if c["change"])
    print(f"  {len(changes)} SKUs, {moved} with a price change, {len(hide)} to hide")
    print(f"  -> {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
